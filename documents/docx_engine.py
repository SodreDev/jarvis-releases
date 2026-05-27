# -*- coding: utf-8 -*-
"""documents.docx_engine — Geração de relatórios Word com tema Stark."""

from __future__ import annotations
import os
import sys
import time
from typing import List, Optional

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .design_system import get_theme, hex_to_rgb, PALETTE


def _base() -> str:
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _out_dir() -> str:
    d = os.path.join(_base(), "data", "documents")
    os.makedirs(d, exist_ok=True)
    return d


def _rgb(hex_color: str) -> RGBColor:
    r, g, b = hex_to_rgb(hex_color)
    return RGBColor(r, g, b)


def _shade(cell, color_hex: str) -> None:
    """Aplica cor de fundo a uma célula de tabela."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex.lstrip("#"))
    tc_pr.append(shd)


def gerar_relatorio(
    titulo: str,
    secoes: List[dict],
    *,
    subtitulo: str = "",
    autor: str = "J.A.R.V.I.S",
    theme_name: str = "stark_light",
    filename: Optional[str] = None,
) -> str:
    """
    Gera um .docx com tema Stark (light por padrão para legibilidade impressa).

    `secoes` = [{"heading": str, "paragraphs": [str,...], "image": path?, "table": [[cell,...],...]?}]
    """
    theme = get_theme(theme_name)
    doc = Document()

    # ── Capa ──
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(titulo)
    title_run.font.name = "Segoe UI"
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = _rgb(theme.accent)

    if subtitulo:
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sp.add_run(subtitulo)
        sr.font.size = Pt(14)
        sr.font.color.rgb = _rgb(PALETTE["text_dim"])

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(f"{autor} • {time.strftime('%d/%m/%Y')}")
    mr.font.size = Pt(10)
    mr.font.color.rgb = _rgb(PALETTE["text_dim"])

    doc.add_paragraph()  # espaço

    # ── Seções ──
    for sec in secoes:
        h = doc.add_heading(sec.get("heading", ""), level=1)
        for run in h.runs:
            run.font.color.rgb = _rgb(theme.accent)
            run.font.name = "Segoe UI"

        for para in sec.get("paragraphs", []) or []:
            p = doc.add_paragraph(para)
            for r in p.runs:
                r.font.name = "Segoe UI"
                r.font.size = Pt(11)

        img = sec.get("image")
        if img and os.path.exists(img):
            try:
                doc.add_picture(img, width=Inches(6))
            except Exception:
                pass

        tbl_data = sec.get("table")
        if tbl_data:
            rows = len(tbl_data)
            cols = max(len(r) for r in tbl_data) if rows else 0
            if rows and cols:
                tbl = doc.add_table(rows=rows, cols=cols)
                tbl.style = "Light Grid Accent 1"
                for i, row in enumerate(tbl_data):
                    for j, val in enumerate(row):
                        cell = tbl.cell(i, j)
                        cell.text = str(val)
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.name = "Segoe UI"
                                r.font.size = Pt(10)
                                if i == 0:
                                    r.font.bold = True
                                    r.font.color.rgb = _rgb("#ffffff")
                        if i == 0:
                            _shade(cell, theme.accent)

        doc.add_paragraph()

    safe = "".join(c if c.isalnum() or c in "-_ " else "_" for c in (filename or titulo))[:60].strip() or "relatorio"
    out = os.path.join(_out_dir(), f"{safe}_{int(time.time())}.docx")
    doc.save(out)
    return out


__all__ = ["gerar_relatorio"]
