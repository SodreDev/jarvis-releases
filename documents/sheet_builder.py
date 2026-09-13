# -*- coding: utf-8 -*-
"""documents.sheet_builder — Planilhas XLSX spec→render (roadmap/13, Fase 2).

O LLM compõe uma SPEC (planilha + abas, cada uma uma tabela e/ou um painel de
KPIs + gráfico). Renderiza um .xlsx PROFISSIONAL via openpyxl, dirigido pelos
design tokens do `design_system`: banda de cabeçalho colorida, linhas zebra,
freeze panes, formatos numéricos (moeda/percent/data), FORMATAÇÃO CONDICIONAL
(escala de cor / data bars / ícones / limiar), GRÁFICOS nativos e auto-largura.
Define page_setup/print_area p/ exportar bonito via `pdf_convert.to_pdf`.

Spec:
  {
    "titulo": str, "tema": str?,
    "abas": [ {
       "nome": str,
       "kpis": [ {valor, titulo, legenda?} ]?,        # painel de KPIs no topo
       "headers": [str], "rows": [[...]],
       "formatos": { "<col nome ou índice>": "moeda|percent|numero|decimal|data|inteiro" }?,
       "condicional": [ {coluna, tipo: "escala|databar|icones|maior|menor", valor?} ]?,
       "grafico": { tipo: "bar|line|pie", titulo?, col_valor: <nome/idx>, col_rotulo: <nome/idx>? }?
    } ]
  }
"""
from __future__ import annotations

import os
import time
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import (
    ColorScaleRule, DataBarRule, IconSetRule, CellIsRule,
)
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

from .design_system import get_theme, build_custom_theme, parse_color_from_text, THEMES, hex_to_rgb


_FMT = {
    "moeda": 'R$ #,##0.00', "currency": 'R$ #,##0.00', "real": 'R$ #,##0.00',
    "percent": '0.0%', "porcentagem": '0.0%', "%": '0.0%',
    "numero": '#,##0', "number": '#,##0', "inteiro": '0',
    "decimal": '#,##0.00', "data": 'DD/MM/YYYY', "date": 'DD/MM/YYYY',
}


def _hex(c: str) -> str:
    return (c or "#000000").lstrip("#").upper()[:6].ljust(6, "0")


def _resolve_theme(tema: Optional[str]):
    if not tema:
        return get_theme("stark_light")
    t = str(tema).strip()
    if t in THEMES:
        return get_theme(t)
    cor = parse_color_from_text(t)
    if cor:
        dark = any(w in t.lower() for w in ("escuro", "dark", "preto", "black"))
        return build_custom_theme(cor, dark=dark)
    return get_theme("stark_light")


def _col_index(ref, headers) -> int:
    """Resolve 'col' (nome do header OU índice 0-based) → índice 0-based."""
    if isinstance(ref, int):
        return ref
    s = str(ref)
    if s.isdigit():
        return int(s)
    low = [str(h).strip().lower() for h in headers]
    try:
        return low.index(s.strip().lower())
    except ValueError:
        return 0


def _auto_width(ws, headers, rows, start_row):
    for j, h in enumerate(headers):
        maxlen = len(str(h))
        for r in rows:
            if j < len(r):
                maxlen = max(maxlen, len(str(r[j])))
        ws.column_dimensions[get_column_letter(j + 1)].width = min(48, max(10, maxlen + 3))


def _render_aba(wb, aba, theme, first):
    nome = (aba.get("nome") or "Planilha")[:31]
    ws = wb.active if first else wb.create_sheet()
    ws.title = nome

    accent = _hex(theme.accent)
    accent_dim = _hex(theme.accent_dim)
    panel = _hex(theme.panel)
    text = _hex(theme.text if theme.is_dark else "#0a0e1a")

    hdr_fill = PatternFill("solid", fgColor=accent_dim)
    hdr_font = Font(bold=True, color="FFFFFF", name="Segoe UI", size=11)
    zebra = PatternFill("solid", fgColor=panel)
    thin = Side(style="thin", color=_hex(theme.grid))
    border = Border(bottom=thin)

    row_cursor = 1

    # ── Painel de KPIs (topo) ────────────────────────────────────────────
    kpis = aba.get("kpis") or []
    if kpis:
        for i, kp in enumerate(kpis[:6]):
            if not isinstance(kp, dict):
                kp = {"valor": str(kp)}
            c = i * 2 + 1
            vcell = ws.cell(row=row_cursor, column=c, value=kp.get("valor", ""))
            vcell.font = Font(bold=True, size=22, color=accent, name="Segoe UI")
            vcell.alignment = Alignment(horizontal="left")
            tcell = ws.cell(row=row_cursor + 1, column=c, value=kp.get("titulo", ""))
            tcell.font = Font(size=11, color=text, name="Segoe UI")
            if kp.get("legenda"):
                lc = ws.cell(row=row_cursor + 2, column=c, value=kp["legenda"])
                lc.font = Font(size=9, color=_hex(theme.grid), name="Segoe UI")
        row_cursor += 4

    headers = aba.get("headers") or []
    rows = aba.get("rows") or []
    if not headers and rows and isinstance(rows[0], list):
        headers = [f"Col {i+1}" for i in range(len(rows[0]))]

    if headers:
        hdr_row = row_cursor
        # Cabeçalho
        for j, h in enumerate(headers):
            cell = ws.cell(row=hdr_row, column=j + 1, value=str(h))
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        # Linhas
        formatos = aba.get("formatos") or {}
        fmt_by_col = {}
        if isinstance(formatos, list):   # forma strict-safe: [{coluna, formato}]
            for it in formatos:
                if isinstance(it, dict) and it.get("coluna") is not None:
                    f = str(it.get("formato", "")).lower()
                    fmt_by_col[_col_index(it["coluna"], headers)] = _FMT.get(f, f)
        elif isinstance(formatos, dict):
            for k, v in formatos.items():
                fmt_by_col[_col_index(k, headers)] = _FMT.get(str(v).lower(), str(v))
        for ri, r in enumerate(rows):
            r = r if isinstance(r, list) else [r]
            excel_row = hdr_row + 1 + ri
            for j in range(len(headers)):
                val = r[j] if j < len(r) else ""
                # tenta numérico p/ formatos numéricos
                if j in fmt_by_col and isinstance(val, str):
                    try:
                        val = float(val.replace("R$", "").replace("%", "").replace(".", "").replace(",", ".").strip())
                    except Exception:
                        pass
                cell = ws.cell(row=excel_row, column=j + 1, value=val)
                cell.font = Font(name="Segoe UI", size=10, color=text)
                cell.border = border
                if j in fmt_by_col:
                    cell.number_format = fmt_by_col[j]
                if ri % 2 == 1:
                    cell.fill = zebra
        last_row = hdr_row + len(rows)
        # Freeze (abaixo do cabeçalho)
        ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
        _auto_width(ws, headers, rows, hdr_row)

        # ── Formatação condicional ───────────────────────────────────────
        for cond in (aba.get("condicional") or []):
            try:
                col = _col_index(cond.get("coluna", 0), headers) + 1
                letter = get_column_letter(col)
                rng = f"{letter}{hdr_row+1}:{letter}{last_row}"
                tipo = (cond.get("tipo") or "escala").lower()
                if tipo in ("escala", "colorscale", "heatmap"):
                    ws.conditional_formatting.add(rng, ColorScaleRule(
                        start_type="min", start_color=_hex(theme.crit if hasattr(theme, "crit") else "#ff4d6d"),
                        mid_type="percentile", mid_value=50, mid_color="FFEB84",
                        end_type="max", end_color="63BE7B"))
                elif tipo in ("databar", "barra"):
                    ws.conditional_formatting.add(rng, DataBarRule(
                        start_type="min", end_type="max", color=accent))
                elif tipo in ("icones", "iconset", "icons"):
                    ws.conditional_formatting.add(rng, IconSetRule(
                        "3TrafficLights1", "percent", [0, 33, 67]))
                elif tipo in ("maior", "acima", "greater"):
                    ws.conditional_formatting.add(rng, CellIsRule(
                        operator="greaterThan", formula=[str(cond.get("valor", 0))],
                        fill=PatternFill("solid", fgColor="C6EFCE"),
                        font=Font(color="006100")))
                elif tipo in ("menor", "abaixo", "less"):
                    ws.conditional_formatting.add(rng, CellIsRule(
                        operator="lessThan", formula=[str(cond.get("valor", 0))],
                        fill=PatternFill("solid", fgColor="FFC7CE"),
                        font=Font(color="9C0006")))
            except Exception as exc:
                print(f"[XLSX] condicional falhou: {exc}")

        # ── Gráfico nativo ───────────────────────────────────────────────
        g = aba.get("grafico")
        if g and rows:
            try:
                _add_chart(ws, g, headers, hdr_row, last_row, theme)
            except Exception as exc:
                print(f"[XLSX] gráfico falhou: {exc}")

        # page setup p/ export PDF limpo
        try:
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.print_area = f"A1:{get_column_letter(len(headers))}{last_row}"
        except Exception:
            pass


def _add_chart(ws, g, headers, hdr_row, last_row, theme):
    tipo = (g.get("tipo") or "bar").lower()
    val_col = _col_index(g.get("col_valor", len(headers) - 1), headers) + 1
    chart = {"line": LineChart, "pie": PieChart}.get(tipo, BarChart)()
    chart.title = g.get("titulo") or ""
    chart.height = 8
    chart.width = 15
    data = Reference(ws, min_col=val_col, min_row=hdr_row, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    if g.get("col_rotulo") is not None:
        lab_col = _col_index(g.get("col_rotulo"), headers) + 1
        cats = Reference(ws, min_col=lab_col, min_row=hdr_row + 1, max_row=last_row)
        chart.set_categories(cats)
    anchor = f"{get_column_letter(len(headers) + 2)}{hdr_row}"
    ws.add_chart(chart, anchor)


def _out_dir() -> str:
    from .pptx_engine import _out_dir as _od
    return _od()


def gerar_planilha(spec: dict, *, filename: Optional[str] = None) -> str:
    """Renderiza a spec → .xlsx. Retorna o caminho do arquivo."""
    if not isinstance(spec, dict):
        raise ValueError("spec deve ser objeto")
    abas = spec.get("abas") or spec.get("sheets") or []
    if not isinstance(abas, list) or not abas:
        # aceita uma planilha simples direto (headers/rows no topo)
        if spec.get("headers") or spec.get("rows"):
            abas = [{"nome": spec.get("titulo", "Planilha"),
                     "headers": spec.get("headers"), "rows": spec.get("rows"),
                     "formatos": spec.get("formatos"), "condicional": spec.get("condicional"),
                     "grafico": spec.get("grafico"), "kpis": spec.get("kpis")}]
        else:
            raise ValueError("spec.abas vazio")
    theme = _resolve_theme(spec.get("tema"))

    wb = Workbook()
    for i, aba in enumerate(abas):
        if isinstance(aba, dict):
            _render_aba(wb, aba, theme, first=(i == 0))

    base = filename or spec.get("titulo") or "planilha"
    safe = "".join(c if c.isalnum() or c in "-_ " else "_" for c in base)[:60].strip() or "planilha"
    out = os.path.join(_out_dir(), f"{safe}_{int(time.time())}.xlsx")
    wb.save(out)
    return out


def capability_manifest() -> dict:
    return {
        "formatos": list({v for v in _FMT}),
        "condicional": ["escala", "databar", "icones", "maior", "menor"],
        "graficos": ["bar", "line", "pie"],
        "estrutura": ("spec={titulo, tema?, abas:[{nome, kpis?[{valor,titulo,legenda?}], "
                      "headers[], rows[[...]], formatos?{col:moeda|percent|numero|decimal|data}, "
                      "condicional?[{coluna,tipo,valor?}], grafico?{tipo,titulo?,col_valor,col_rotulo?}}]}"),
        "regras": ("Use 'formatos' p/ moeda/percent/data por coluna; 'condicional' p/ realçar "
                   "(escala de cor em métricas, maior/menor p/ metas); 'grafico' p/ visualizar; "
                   "'kpis' p/ um painel-resumo no topo. Múltiplas 'abas' p/ seções."),
    }


__all__ = ["gerar_planilha", "capability_manifest"]
